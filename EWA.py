#!/usr/bin/env python3
"""
SAP EWA (Early Watch Alert) Report Analyzer
Parses PDF/HTML EWA reports and generates a formatted Excel action item tracker.

Usage:
    python ewa_analyzer.py --file "C:/path/to/EWA_Report.pdf"
    python ewa_analyzer.py --file "C:/path/to/EWA_Report.html" --output "C:/out/result.xlsx" --sid PRD
"""

import argparse
import os
import re
import sys
from datetime import datetime
from pathlib import Path

# ─────────────────────────────────────────────
# Dependency check
# ─────────────────────────────────────────────
def check_dependencies():
    missing = []
    for pkg, imp in [("pdfplumber", "pdfplumber"), ("beautifulsoup4", "bs4"), ("openpyxl", "openpyxl")]:
        try:
            __import__(imp)
        except ImportError:
            missing.append(pkg)
    if missing:
        print("\n[ERROR] Missing packages. Install them with:")
        print(f"        pip install {' '.join(missing)}\n")
        sys.exit(1)

check_dependencies()

import pdfplumber
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side, GradientFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


# ─────────────────────────────────────────────
# Constants / Mapping tables
# ─────────────────────────────────────────────

CATEGORY_PATTERNS = {
    "Performance":          [r"response time", r"throughput", r"dialog.*time", r"workload", r"cpu.*utiliz"],
    "Database":             [r"database", r"tablespace", r"table growth", r"\bSQL\b", r"index", r"db statistics", r"oracle", r"hana", r"db size"],
    "Security":             [r"security", r"authoriz", r"password", r"SAP_ALL", r"SAP\*", r"\bDDIC\b", r"firefighter", r"emergency user"],
    "Memory / Buffers":     [r"memory", r"buffer", r"paging", r"swapping", r"extended memory", r"heap", r"roll area", r"EM\b"],
    "ABAP Runtime":         [r"short dump", r"ABAP", r"runtime error", r"\bdump\b", r"ST22"],
    "Support Package":      [r"support package", r"\bSP\s+level\b", r"patch level", r"kernel.*patch", r"upgrade"],
    "Background Processing":[r"background job", r"batch job", r"SM37", r"job.*fail", r"job.*cancel", r"job.*abort"],
    "Update Processing":    [r"update.*error", r"update.*terminat", r"SM13", r"update.*fail"],
    "Enqueue / Locks":      [r"enqueue", r"\block\b", r"deadlock", r"SM12"],
    "Output Management":    [r"spool", r"output.*queue", r"print.*fail", r"SP01"],
    "Transport Management": [r"transport", r"STMS", r"import queue", r"\bCTS\b", r"transport.*fail"],
    "Availability":         [r"availability", r"uptime", r"downtime", r"system.*restart", r"instance.*restart"],
    "Backup & Recovery":    [r"backup.*fail", r"no.*backup", r"backup.*miss", r"restore", r"archive log"],
    "IDoc / Interface":     [r"IDoc", r"interface.*error", r"\bRFC\b.*error", r"middleware", r"\bPI\b", r"\bPO\b"],
    "Workload":             [r"work process", r"dialog.*WP", r"background.*WP", r"RFC.*WP"],
    "System Configuration": [r"profile parameter", r"instance profile", r"configuration", r"parameter.*change"],
}

PRIORITY_PATTERNS = {
    "HIGH":   [r"\bRED\b", r"\bcritical\b", r"immediately", r"\bsevere\b", r"\burgent\b",
               r"must be corrected", r"action required", r"exceeded.*limit", r"not acceptable",
               r"serious.*issue", r"needs.*immediate", r"fails?\b"],
    "MEDIUM": [r"\bYELLOW\b", r"\bwarning\b", r"\battention\b", r"\bmonitor\b",
               r"suboptimal", r"should be", r"recommended", r"not optimal", r"could be improved",
               r"consider.*changing", r"review.*recommended"],
    "LOW":    [r"\bGREEN\b", r"\binformation\b", r"\bnote\b", r"\bOK\b",
               r"acceptable", r"no action required", r"for information", r"within.*limit",
               r"within.*range", r"satisfactory"],
}

# Pre-built well-known EWA checks ─ always emit if detected in text
KNOWN_CHECKS = [
    # ── Security ──────────────────────────────────────────────────
    {
        "pattern": r"SAP\*.*?password|default.*?password.*?SAP\*|password.*?SAP\*.*?default",
        "category": "Security", "base_priority": "HIGH",
        "description": "Default SAP* Password Not Changed",
        "recommendation": (
            "Change the default password for SAP*, DDIC, and EARLYWATCH users immediately. "
            "Use transaction SU01. Locking SAP* in all clients is also recommended."
        ),
    },
    {
        "pattern": r"SAP_ALL.*?assign|assign.*?SAP_ALL|users.*?SAP_ALL",
        "category": "Security", "base_priority": "HIGH",
        "description": "SAP_ALL Profile Assigned to Business Users",
        "recommendation": (
            "Remove SAP_ALL from non-emergency users. Implement role-based authorization. "
            "Use SU01/SUIM to identify and remediate affected users."
        ),
    },
    {
        "pattern": r"inactive.*?user|user.*?inactive|locked.*?user.*?never.*?log",
        "category": "Security", "base_priority": "MEDIUM",
        "description": "Inactive / Never-Logged-In Users Exist",
        "recommendation": (
            "Run SUIM to list users inactive > 90 days. Lock or delete stale accounts "
            "to reduce the attack surface."
        ),
    },
    # ── Performance ───────────────────────────────────────────────
    {
        "pattern": r"dialog.*?response.*?time|avg.*?response.*?time.*?high",
        "category": "Performance", "base_priority": "HIGH",
        "description": "High Dialog Response Time Detected",
        "recommendation": (
            "Use SM66/SM50 to identify long-running work processes. "
            "Analyse top programs in ST05/SE30. Review work process allocation."
        ),
    },
    {
        "pattern": r"database.*?response.*?time.*?high|DB.*?response.*?time.*?exceed",
        "category": "Database", "base_priority": "HIGH",
        "description": "High Database Response Time",
        "recommendation": (
            "Run DB02/ST05 to find expensive SQL. Update DB statistics. "
            "Check for missing indexes and evaluate query execution plans."
        ),
    },
    {
        "pattern": r"RFC.*?response.*?time|RFC.*?slow",
        "category": "Performance", "base_priority": "MEDIUM",
        "description": "High RFC Response Time",
        "recommendation": (
            "Check SM59 RFC destinations. Analyse RFC bottlenecks in SM66. "
            "Review network latency between systems."
        ),
    },
    # ── Database ──────────────────────────────────────────────────
    {
        "pattern": r"tablespace.*?full|tablespace.*?critical|disk.*?space.*?critical|tablespace.*?9[0-9]%",
        "category": "Database", "base_priority": "HIGH",
        "description": "Tablespace / Disk Space Critical (>90%)",
        "recommendation": (
            "Immediately add datafiles to critical tablespaces in DB02. "
            "Implement data archiving (SARA) and review table growth in DB02."
        ),
    },
    {
        "pattern": r"table.*?growth|large.*?table|growing.*?table",
        "category": "Database", "base_priority": "MEDIUM",
        "description": "Excessive Table Growth Detected",
        "recommendation": (
            "Identify top tables by size in DB02. Implement archiving strategy via SARA. "
            "Consider partitioning large tables."
        ),
    },
    {
        "pattern": r"db.*?statistic.*?outdated|outdated.*?statistic|missing.*?statistic",
        "category": "Database", "base_priority": "MEDIUM",
        "description": "Outdated / Missing Database Statistics",
        "recommendation": (
            "Schedule regular statistics update via DB13 (Oracle) or HANA Studio. "
            "Outdated statistics cause suboptimal query execution plans."
        ),
    },
    # ── Memory / Buffers ──────────────────────────────────────────
    {
        "pattern": r"extended memory.*?exhaust|EM.*?exhaust|memory.*?swap|paging.*?occur",
        "category": "Memory / Buffers", "base_priority": "HIGH",
        "description": "Extended Memory Exhaustion / OS Swapping",
        "recommendation": (
            "Increase em/initial_size_MB and ztta/roll_extension. "
            "Check RZ20 memory alerts. Analyse top memory consumers with OS07."
        ),
    },
    {
        "pattern": r"buffer.*?hit.*?ratio.*?low|low.*?buffer.*?hit|buffer.*?quality.*?poor",
        "category": "Memory / Buffers", "base_priority": "MEDIUM",
        "description": "Low Buffer Hit Ratio",
        "recommendation": (
            "Increase buffer sizes: zcsa/table_buffer_area, rsdb/ntab/buffersize, "
            "abap/buffersize. Monitor with ST02."
        ),
    },
    # ── ABAP Runtime ──────────────────────────────────────────────
    {
        "pattern": r"short dump|ABAP.*?dump|dump.*?frequent|high.*?dump.*?rate",
        "category": "ABAP Runtime", "base_priority": "MEDIUM",
        "description": "Frequent ABAP Short Dumps (Runtime Errors)",
        "recommendation": (
            "Analyse top dump classes in ST22. Apply relevant SAP Notes. "
            "Fix custom ABAP code causing dumps. Set dump retention appropriately."
        ),
    },
    # ── Support Package ───────────────────────────────────────────
    {
        "pattern": r"support package.*?behind|SP.*?level.*?old|outdated.*?support package|kernel.*?outdated",
        "category": "Support Package", "base_priority": "MEDIUM",
        "description": "Support Package / Kernel Level Outdated",
        "recommendation": (
            "Plan SP stack upgrade. Download latest SP stack from SAP Maintenance Planner. "
            "Apply critical SAP Notes before full SP upgrade."
        ),
    },
    # ── Background Processing ─────────────────────────────────────
    {
        "pattern": r"background job.*?fail|batch.*?cancel|job.*?abort|critical.*?job.*?fail",
        "category": "Background Processing", "base_priority": "MEDIUM",
        "description": "Background / Batch Job Failures",
        "recommendation": (
            "Review SM37 for cancelled/failed jobs. Identify root cause. "
            "Schedule reruns and implement job monitoring alerts."
        ),
    },
    # ── Update Processing ─────────────────────────────────────────
    {
        "pattern": r"update.*?error|update.*?terminat|SM13.*?error|update.*?fail",
        "category": "Update Processing", "base_priority": "HIGH",
        "description": "Update Errors / Terminations in SM13",
        "recommendation": (
            "Open SM13 and analyse update terminations. Reprocess valid records, "
            "cancel irrelevant ones. Identify root cause and fix affected programs."
        ),
    },
    # ── Backup & Recovery ─────────────────────────────────────────
    {
        "pattern": r"backup.*?fail|no.*?backup|backup.*?not.*?run|backup.*?miss",
        "category": "Backup & Recovery", "base_priority": "HIGH",
        "description": "Database Backup Failure / Backup Not Running",
        "recommendation": (
            "Verify backup configuration via DB13. Ensure daily full backup and "
            "regular log backup. Test restore procedure quarterly."
        ),
    },
    # ── IDoc / Interface ──────────────────────────────────────────
    {
        "pattern": r"IDoc.*?error|IDoc.*?fail|interface.*?error|RFC.*?fail|RFC.*?connection",
        "category": "IDoc / Interface", "base_priority": "MEDIUM",
        "description": "IDoc / RFC / Interface Errors",
        "recommendation": (
            "Check WE02/BD87 for IDoc errors. Verify SM59 RFC destinations. "
            "Reprocess failed IDocs and resolve connectivity issues."
        ),
    },
    # ── Availability ──────────────────────────────────────────────
    {
        "pattern": r"availability.*?below|downtime.*?unplanned|system.*?restart.*?unexpected",
        "category": "Availability", "base_priority": "HIGH",
        "description": "System Availability Below Target / Unplanned Downtime",
        "recommendation": (
            "Analyse system logs (SM21) for crash root cause. "
            "Review core dumps, OS logs. Consider HA configuration if not already in place."
        ),
    },
    # ── Workload ──────────────────────────────────────────────────
    {
        "pattern": r"work process.*?util|WP.*?utiliz.*?high|all.*?WP.*?busy",
        "category": "Workload", "base_priority": "HIGH",
        "description": "High Work Process Utilization (WP Bottleneck)",
        "recommendation": (
            "Review work process allocation in RZ10. Consider adding dialog/background "
            "WPs or additional application servers to distribute load."
        ),
    },
    # ── Enqueue / Locks ───────────────────────────────────────────
    {
        "pattern": r"lock.*?conflict|enqueue.*?error|deadlock|lock.*?overflow",
        "category": "Enqueue / Locks", "base_priority": "MEDIUM",
        "description": "Lock / Enqueue Conflicts or Overflow",
        "recommendation": (
            "Check SM12 for old lock entries. Review enqueue server sizing "
            "(enque/table_size). Analyse programs causing long-held locks."
        ),
    },
    # ── System Configuration ──────────────────────────────────────
    {
        "pattern": r"profile parameter.*?incorrect|parameter.*?not.*?optimal|rz10.*?change",
        "category": "System Configuration", "base_priority": "MEDIUM",
        "description": "Non-Optimal Profile Parameters Detected",
        "recommendation": (
            "Review SAP recommendations in RZ10. Compare with SAP Note 103747 "
            "(parameter recommendations). Apply changes in maintenance window."
        ),
    },
]


# ─────────────────────────────────────────────
# Parsers
# ─────────────────────────────────────────────

class EWAPDFParser:
    def __init__(self, filepath):
        self.filepath = filepath
        self.full_text = ""
        self.page_count = 0

    def parse(self):
        print(f"  [PDF] Parsing: {self.filepath}")
        with pdfplumber.open(self.filepath) as pdf:
            self.page_count = len(pdf.pages)
            for page in pdf.pages:
                self.full_text += (page.extract_text() or "") + "\n"
        print(f"  [PDF] {self.page_count} pages, {len(self.full_text):,} chars extracted")
        return self.full_text

    def extract_date(self):
        return _extract_date_from_text(self.full_text[:4000])


class EWAHTMLParser:
    def __init__(self, filepath):
        self.filepath = filepath
        self.full_text = ""

    def parse(self):
        print(f"  [HTML] Parsing: {self.filepath}")
        with open(self.filepath, "r", encoding="utf-8", errors="ignore") as fh:
            soup = BeautifulSoup(fh.read(), "html.parser")
        self.full_text = soup.get_text(separator="\n")
        print(f"  [HTML] {len(self.full_text):,} chars extracted")
        return self.full_text

    def extract_date(self):
        return _extract_date_from_text(self.full_text[:4000])


def _extract_date_from_text(text):
    patterns = [
        r"Report\s+Date[:\s]+(\d{1,2}[./\-]\d{1,2}[./\-]\d{2,4})",
        r"Generated\s+on[:\s]+(\d{1,2}[./\-]\d{1,2}[./\-]\d{2,4})",
        r"EarlyWatch\s+Alert.*?(\d{1,2}[./\-]\d{1,2}[./\-]\d{2,4})",
        r"Date[:\s]+(\d{1,2}[./\-]\d{1,2}[./\-]\d{2,4})",
        r"(\d{1,2}\.\d{1,2}\.\d{4})",
        r"(\d{4}-\d{2}-\d{2})",
        r"(\w+ \d{1,2},?\s+\d{4})",
    ]
    for p in patterns:
        m = re.search(p, text, re.IGNORECASE)
        if m:
            return m.group(1).strip()
    return datetime.now().strftime("%d.%m.%Y")


# ─────────────────────────────────────────────
# Finding Extractor
# ─────────────────────────────────────────────

def _determine_category(snippet):
    low = snippet.lower()
    for cat, pats in CATEGORY_PATTERNS.items():
        if any(re.search(p, low, re.IGNORECASE) for p in pats):
            return cat
    return "System Configuration"


def _determine_priority(snippet):
    for priority, pats in PRIORITY_PATTERNS.items():
        if any(re.search(p, snippet, re.IGNORECASE) for p in pats):
            return priority
    return "MEDIUM"


def extract_findings(text):
    findings = []
    seen_keys = set()

    # ── Pass 1: known pre-built checks ───────────────────────────
    for check in KNOWN_CHECKS:
        m = re.search(check["pattern"], text, re.IGNORECASE)
        if m:
            context = text[max(0, m.start() - 150): m.end() + 250]
            detected_priority = _determine_priority(context)
            # Never downgrade a HIGH base check to LOW
            priority = detected_priority
            if check["base_priority"] == "HIGH" and detected_priority == "LOW":
                priority = "HIGH"

            key = check["description"][:40].lower()
            if key not in seen_keys:
                seen_keys.add(key)
                findings.append({
                    "category":       check["category"],
                    "priority":       priority,
                    "description":    check["description"],
                    "recommendation": check["recommendation"],
                    "source":         "EWA Check",
                })

    # ── Pass 2: RED / YELLOW / GREEN inline ratings ───────────────
    rating_block = re.compile(
        r"(RED|YELLOW|GREEN)\s*[:\-–]\s*(.{20,300}?)(?=\n\n|RED\s*[:\-]|YELLOW\s*[:\-]|GREEN\s*[:\-]|\Z)",
        re.IGNORECASE | re.DOTALL,
    )
    for m in rating_block.finditer(text):
        rating   = m.group(1).upper()
        content  = re.sub(r"\s+", " ", m.group(2)).strip()
        priority = {"RED": "HIGH", "YELLOW": "MEDIUM", "GREEN": "LOW"}.get(rating, "MEDIUM")
        category = _determine_category(content)
        rec      = _pull_recommendation(content)
        key      = content[:40].lower()
        if key not in seen_keys and len(content) > 30:
            seen_keys.add(key)
            findings.append({
                "category":       category,
                "priority":       priority,
                "description":    content[:220],
                "recommendation": rec,
                "source":         f"Rating:{rating}",
            })

    # ── Pass 3: Recommendation / Action blocks ────────────────────
    rec_block = re.compile(
        r"(?:Recommendation[s]?|Suggested Action[s]?|Action Required)[:\s]+(.{30,500}?)(?=\n\n|\n[A-Z][A-Z]|\Z)",
        re.IGNORECASE | re.DOTALL,
    )
    for m in rec_block.finditer(text):
        rec_text = re.sub(r"\s+", " ", m.group(1)).strip()
        ctx_start = max(0, m.start() - 350)
        context   = text[ctx_start: m.start()]
        priority  = _determine_priority(context + rec_text)
        category  = _determine_category(context + rec_text)
        desc_lines = [l.strip() for l in context.split("\n") if len(l.strip()) > 25]
        description = desc_lines[-1][:220] if desc_lines else context[-120:].strip()
        key = description[:40].lower()
        if key not in seen_keys and len(description) > 25:
            seen_keys.add(key)
            findings.append({
                "category":       category,
                "priority":       priority,
                "description":    description,
                "recommendation": rec_text[:300],
                "source":         "Recommendation block",
            })

    print(f"  Extracted {len(findings)} unique findings")
    return findings


def _pull_recommendation(text):
    for pat in [
        r"recommend[ation]*s?[:\s]+(.{20,300})",
        r"action[s]?\s+required[:\s]+(.{20,300})",
        r"should[:\s]+(.{20,200})",
        r"please[:\s]+(.{20,200})",
    ]:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            return re.sub(r"\s+", " ", m.group(1)).strip()
    return "Review finding and implement SAP recommendation."


# ─────────────────────────────────────────────
# Excel Generator
# ─────────────────────────────────────────────

PALETTE = {
    "header_bg":   "1F4E79",
    "header_fg":   "FFFFFF",
    "high_bg":     "FFCCCC",
    "medium_bg":   "FFEB99",
    "low_bg":      "CCFFCC",
    "high_txt":    "CC0000",
    "medium_txt":  "7F4F00",
    "low_txt":     "1A6B1A",
    "alt_row":     "EEF4FB",
    "open_bg":     "FFE699",
    "closed_bg":   "C6EFCE",
    "inprog_bg":   "BDD7EE",
    "section_bg":  "2E75B6",
    "border":      "B8CCE4",
}

def _fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

def _border():
    thin = Side(style="thin", color=PALETTE["border"])
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def _font(bold=False, size=10, color="000000", name="Calibri"):
    return Font(name=name, bold=bold, size=size, color=color)

def _align(h="left", v="top", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def build_excel(findings, report_date, source_file, output_path, sid=""):
    wb = openpyxl.Workbook()

    # ── Sheet 1: Action Items ─────────────────────────────────────
    ws = wb.active
    ws.title = "Action Items"

    COLUMNS = [
        ("No.",                5),
        ("EWA Report Date",   14),
        ("System / SID",      12),
        ("Priority",          10),
        ("Type / Category",   22),
        ("Finding / Description", 48),
        ("SAP Recommendation",    48),
        ("Current Status",    13),
        ("Analysis & Action Remark", 42),
        ("Assigned To",       16),
        ("Target Date",       14),
        ("Completed Date",    14),
        ("Remarks",           22),
    ]

    # Header row
    for col, (label, width) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=col, value=label)
        c.fill    = _fill(PALETTE["header_bg"])
        c.font    = _font(bold=True, size=11, color=PALETTE["header_fg"])
        c.alignment = _align("center", "center")
        c.border  = _border()
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32

    # Data rows
    priority_order = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
    sorted_findings = sorted(findings, key=lambda x: priority_order.get(x.get("priority", "MEDIUM"), 1))

    status_dv = DataValidation(type="list", formula1='"Open,In Progress,Closed,N/A,Deferred"', allow_blank=True)
    ws.add_data_validation(status_dv)

    for idx, f in enumerate(sorted_findings, 1):
        row   = idx + 1
        pri   = f.get("priority", "MEDIUM")
        alt   = (idx % 2 == 0)
        row_bg = PALETTE["alt_row"] if alt else "FFFFFF"

        row_data = [
            idx,
            report_date,
            sid or "",
            pri,
            f.get("category", "General"),
            f.get("description", ""),
            f.get("recommendation", ""),
            "Open",
            "",   # Analysis & Action Remark
            "",   # Assigned To
            "",   # Target Date
            "",   # Completed Date
            "",   # Remarks
        ]

        for col, val in enumerate(row_data, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border    = _border()
            c.font      = _font(size=10)

            if col == 1:    # No.
                c.alignment = _align("center", "center", False)
                c.fill = _fill(row_bg)
            elif col == 2:  # Date
                c.alignment = _align("center", "center", False)
                c.fill = _fill(row_bg)
            elif col == 3:  # SID
                c.alignment = _align("center", "center", False)
                c.fill = _fill(row_bg)
            elif col == 4:  # Priority
                bg  = {"HIGH": PALETTE["high_bg"], "MEDIUM": PALETTE["medium_bg"], "LOW": PALETTE["low_bg"]}.get(pri, "FFFFFF")
                txt = {"HIGH": PALETTE["high_txt"], "MEDIUM": PALETTE["medium_txt"], "LOW": PALETTE["low_txt"]}.get(pri, "000000")
                c.fill      = _fill(bg)
                c.font      = _font(bold=True, size=10, color=txt)
                c.alignment = _align("center", "center", False)
            elif col == 8:  # Status
                c.fill      = _fill(PALETTE["open_bg"])
                c.font      = _font(bold=True, size=10)
                c.alignment = _align("center", "center", False)
                status_dv.add(c)
            elif col == 9:  # Remark
                c.fill      = _fill("FFFFF0")
                c.alignment = _align("left", "top", True)
            else:
                c.fill      = _fill(row_bg)
                c.alignment = _align("left", "top", True)

        ws.row_dimensions[row].height = 48

    # Auto-filter + freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(sorted_findings)+1}"
    ws.freeze_panes   = "A2"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage   = True
    ws.page_setup.fitToWidth  = 1

    # ── Sheet 2: Summary Dashboard ────────────────────────────────
    ws2 = wb.create_sheet("Summary")
    ws2.sheet_view.showGridLines = False

    # Title banner
    ws2.merge_cells("A1:G1")
    t = ws2["A1"]
    t.value     = "SAP EWA Report — Analysis Summary"
    t.fill      = _fill(PALETTE["header_bg"])
    t.font      = _font(bold=True, size=15, color="FFFFFF")
    t.alignment = _align("center", "center", False)
    ws2.row_dimensions[1].height = 40

    # Report meta
    meta = [
        ("EWA Report Date",  report_date),
        ("Source File",      os.path.basename(source_file)),
        ("System / SID",     sid or "—"),
        ("Analysis Run At",  datetime.now().strftime("%Y-%m-%d  %H:%M")),
        ("Total Findings",   len(findings)),
    ]
    for i, (label, val) in enumerate(meta, 3):
        ws2.cell(row=i, column=1, value=label).font = _font(bold=True, size=11)
        ws2.cell(row=i, column=2, value=val).font   = _font(size=11)

    # Priority breakdown
    _write_section_header(ws2, 9, "PRIORITY BREAKDOWN")
    for col, h in enumerate(["Priority", "Count", "% of Total"], 1):
        c = ws2.cell(row=10, column=col, value=h)
        c.fill = _fill("BDD7EE"); c.font = _font(bold=True, size=10); c.border = _border()
        c.alignment = _align("center", "center", False)

    counts = {"HIGH": 0, "MEDIUM": 0, "LOW": 0}
    for f in findings:
        counts[f.get("priority", "MEDIUM")] = counts.get(f.get("priority", "MEDIUM"), 0) + 1

    total = len(findings) or 1
    pri_colors = {"HIGH": PALETTE["high_bg"], "MEDIUM": PALETTE["medium_bg"], "LOW": PALETTE["low_bg"]}
    for i, (p, cnt) in enumerate(counts.items(), 11):
        ws2.cell(row=i, column=1, value=p).fill  = _fill(pri_colors[p])
        ws2.cell(row=i, column=2, value=cnt)
        ws2.cell(row=i, column=3, value=f"{cnt/total*100:.1f}%")
        for col in range(1, 4):
            ws2.cell(row=i, column=col).border    = _border()
            ws2.cell(row=i, column=col).alignment = _align("center", "center", False)
            ws2.cell(row=i, column=col).font      = _font(size=10)

    # Category breakdown
    cats = {}
    for f in findings:
        cats[f.get("category", "General")] = cats.get(f.get("category", "General"), 0) + 1
    sorted_cats = sorted(cats.items(), key=lambda x: x[1], reverse=True)

    start_row = 16
    _write_section_header(ws2, start_row, "CATEGORY BREAKDOWN")
    for col, h in enumerate(["Category", "Count", "% of Total"], 1):
        c = ws2.cell(row=start_row+1, column=col, value=h)
        c.fill = _fill("BDD7EE"); c.font = _font(bold=True, size=10); c.border = _border()
        c.alignment = _align("center", "center", False)

    for i, (cat, cnt) in enumerate(sorted_cats, start_row+2):
        bg = PALETTE["alt_row"] if i % 2 == 0 else "FFFFFF"
        ws2.cell(row=i, column=1, value=cat).fill  = _fill(bg)
        ws2.cell(row=i, column=2, value=cnt).fill  = _fill(bg)
        ws2.cell(row=i, column=3, value=f"{cnt/total*100:.1f}%").fill = _fill(bg)
        for col in range(1, 4):
            ws2.cell(row=i, column=col).border    = _border()
            ws2.cell(row=i, column=col).alignment = _align("center", "center", False)
            ws2.cell(row=i, column=col).font      = _font(size=10)

    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 10
    ws2.column_dimensions["C"].width = 14

    # ── Sheet 3: Standard EWA Checklist ──────────────────────────
    ws3 = wb.create_sheet("EWA Checklist")
    chk_headers = ["#", "Check Area", "Check Item", "Expected Threshold", "Actual Finding", "Status", "Action Required"]
    chk_widths  = [4, 22, 36, 28, 32, 12, 36]

    for col, (h, w) in enumerate(zip(chk_headers, chk_widths), 1):
        c = ws3.cell(row=1, column=col, value=h)
        c.fill = _fill(PALETTE["header_bg"]); c.font = _font(bold=True, size=11, color="FFFFFF")
        c.alignment = _align("center", "center"); c.border = _border()
        ws3.column_dimensions[get_column_letter(col)].width = w
    ws3.row_dimensions[1].height = 30

    checklist = [
        ("Security",    "SAP* default password changed",        "Changed",           "","Check",""),
        ("Security",    "DDIC default password changed",         "Changed",           "","Check",""),
        ("Security",    "EARLYWATCH password changed",           "Changed",           "","Check",""),
        ("Security",    "No user assigned SAP_ALL in PRD",       "0 users",           "","Check",""),
        ("Security",    "Password policy active",                "Minimum complexity","","Check",""),
        ("Security",    "Inactive users locked/deleted",         "0 inactive > 90d",  "","Check",""),
        ("Performance", "Dialog avg response time",              "< 1,000 ms",        "","Check",""),
        ("Performance", "DB avg response time",                  "< 40 ms",           "","Check",""),
        ("Performance", "CPU utilization (avg)",                 "< 70%",             "","Check",""),
        ("Performance", "Memory paging / swapping",              "None",              "","Check",""),
        ("Database",    "Tablespace utilization",                "< 85%",             "","Check",""),
        ("Database",    "DB statistics current",                 "Updated < 7 days",  "","Check",""),
        ("Database",    "Daily DB backup successful",            "Green",             "","Check",""),
        ("Database",    "No missing or invalid indexes",         "0",                 "","Check",""),
        ("Memory",      "Extended memory not exhausted",         "No exhaustion",     "","Check",""),
        ("Memory",      "Buffer hit ratios (table, NTAB, prog)", "> 95%",             "","Check",""),
        ("ABAP",        "Short dump count (last 7 days)",        "< 100",             "","Check",""),
        ("Updates",     "Update terminations in SM13",           "0",                 "","Check",""),
        ("Batch Jobs",  "Failed critical background jobs",       "0",                 "","Check",""),
        ("Availability","System availability",                   "> 99.5%",           "","Check",""),
        ("Sp. Package", "SP stack level",                        "Within 2 SP",       "","Check",""),
        ("Sp. Package", "Kernel patch level",                    "< 6 months old",    "","Check",""),
        ("Interfaces",  "IDoc error rate",                       "< 1%",              "","Check",""),
        ("Interfaces",  "RFC connection errors",                 "0 critical",        "","Check",""),
    ]

    for i, row_data in enumerate(checklist, 2):
        bg = PALETTE["alt_row"] if i % 2 == 0 else "FFFFFF"
        ws3.cell(row=i, column=1, value=i-1)
        for col, val in enumerate(row_data, 2):
            c = ws3.cell(row=i, column=col, value=val)
            c.fill = _fill(bg); c.border = _border()
            c.font = _font(size=10); c.alignment = _align("left", "center", False)
        ws3.row_dimensions[i].height = 18

    ws3.auto_filter.ref = f"A1:{get_column_letter(len(chk_headers))}1"
    ws3.freeze_panes = "A2"

    # ── Save ──────────────────────────────────────────────────────
    wb.save(output_path)
    return output_path


def _write_section_header(ws, row, title):
    ws.merge_cells(f"A{row}:C{row}")
    c = ws.cell(row=row, column=1, value=title)
    c.fill      = _fill(PALETTE["section_bg"])
    c.font      = _font(bold=True, size=11, color="FFFFFF")
    c.alignment = _align("center", "center", False)
    ws.row_dimensions[row].height = 22


# ─────────────────────────────────────────────
# Orchestrator
# ─────────────────────────────────────────────

def analyze(filepath, output_path=None, sid=""):
    filepath = Path(filepath)

    if not filepath.exists():
        print(f"\n[ERROR] File not found: {filepath}")
        sys.exit(1)

    ext = filepath.suffix.lower()
    if ext == ".pdf":
        parser = EWAPDFParser(str(filepath))
    elif ext in (".html", ".htm"):
        parser = EWAHTMLParser(str(filepath))
    else:
        print(f"\n[ERROR] Unsupported file type '{ext}'. Supported: .pdf  .html  .htm")
        sys.exit(1)

    print("\n" + "="*62)
    print("  SAP EWA Report Analyzer")
    print("="*62)
    print(f"  File : {filepath}")

    text        = parser.parse()
    report_date = parser.extract_date()
    print(f"  Date : {report_date}")
    if sid:
        print(f"  SID  : {sid}")

    findings = extract_findings(text)

    if not findings:
        print("\n  [WARN] No findings detected — creating blank 20-row template.")
        findings = [{
            "category": "", "priority": "MEDIUM",
            "description": "", "recommendation": "", "source": "Manual"
        }] * 20

    # Sort HIGH → MEDIUM → LOW
    order = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
    findings.sort(key=lambda x: order.get(x.get("priority", "MEDIUM"), 1))

    if not output_path:
        ts = datetime.now().strftime("%Y%m%d_%H%M")
        output_path = filepath.parent / f"{filepath.stem}_EWA_Analysis_{ts}.xlsx"

    print(f"\n  Generating Excel...")
    saved = build_excel(findings, report_date, str(filepath), str(output_path), sid)

    high   = sum(1 for f in findings if f.get("priority") == "HIGH")
    medium = sum(1 for f in findings if f.get("priority") == "MEDIUM")
    low    = sum(1 for f in findings if f.get("priority") == "LOW")

    print(f"\n{'='*62}")
    print(f"  ANALYSIS COMPLETE")
    print(f"  Total Findings : {len(findings)}")
    print(f"  HIGH  (Red)    : {high}")
    print(f"  MEDIUM (Yellow): {medium}")
    print(f"  LOW   (Green)  : {low}")
    print(f"  Output Excel   : {saved}")
    print(f"{'='*62}\n")
    return saved


# ─────────────────────────────────────────────
# CLI entry point
# ─────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(
        description="SAP EWA Report Analyzer — PDF/HTML → Excel Action Tracker",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python ewa_analyzer.py --file "C:/EWA/report.pdf"
  python ewa_analyzer.py --file "C:/EWA/report.html" --sid PRD
  python ewa_analyzer.py --file "C:/EWA/report.pdf"  --output "C:/Output/EWA_PRD.xlsx" --sid PRD
        """,
    )
    ap.add_argument("--file",   required=True, help="Path to EWA report (.pdf or .html/.htm)")
    ap.add_argument("--output", default=None,  help="Output Excel path (auto-generated if omitted)")
    ap.add_argument("--sid",    default="",    help="SAP System ID, e.g. PRD")
    args = ap.parse_args()
    analyze(args.file, args.output, args.sid)


if __name__ == "__main__":
    main()
